#!/usr/bin/env python3
"""Gerador da Apresentação Executiva MiniHub Car em PowerPoint (.pptx)."""

import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Cores do Tema Automotivo Premium (Dark Mode)
BG_DARK = RGBColor(15, 23, 42)        # #0f172a (Slate 900)
CARD_BG = RGBColor(30, 41, 59)        # #1e293b (Slate 800)
CARD_BORDER = RGBColor(51, 65, 85)    # #334155 (Slate 700)
ACCENT_BLUE = RGBColor(14, 165, 233)  # #0ea5e9 (Sky 500)
ACCENT_ORANGE = RGBColor(249, 115, 22) # #f97316 (Orange 500)
ACCENT_GREEN = RGBColor(16, 185, 129) # #10b981 (Emerald 500)
ACCENT_PURPLE = RGBColor(168, 85, 247) # #a855f7 (Purple 500)
TEXT_WHITE = RGBColor(255, 255, 255)
TEXT_MUTED = RGBColor(148, 163, 184)  # #94a3b8 (Slate 400)
TEXT_LIGHT = RGBColor(226, 232, 240)  # #e2e8f0 (Slate 200)

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"


def create_blank_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_DARK
    bg.line.color.rgb = BG_DARK
    return slide


def add_header(slide, tag: str, title: str, subtitle: str = ""):
    # Tag / Categoria
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
    tf = tag_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = tag.upper()
    p.font.name = FONT_HEADING
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    # Título Principal
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE

    # Subtítulo (se houver)
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.7), Inches(0.5))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = TEXT_MUTED


def add_card(slide, x, y, w, h, title: str, text: str, accent_color=ACCENT_BLUE, badge: str = ""):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = CARD_BORDER
    card.line.width = Pt(1.2)

    # Faixa lateral ou badge de cor
    strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Pt(6), h)
    strip.fill.solid()
    strip.fill.fore_color.rgb = accent_color
    strip.line.fill.background()

    tb = slide.shapes.add_textbox(x + Inches(0.2), y + Inches(0.15), w - Inches(0.35), h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True

    if badge:
        p_badge = tf.paragraphs[0]
        p_badge.text = badge.upper()
        p_badge.font.name = FONT_HEADING
        p_badge.font.size = Pt(9)
        p_badge.font.bold = True
        p_badge.font.color.rgb = accent_color
        p_title = tf.add_paragraph()
    else:
        p_title = tf.paragraphs[0]

    p_title.text = title
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(15)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE
    p_title.space_after = Pt(6)

    p_body = tf.add_paragraph()
    p_body.text = text
    p_body.font.name = FONT_BODY
    p_body.font.size = Pt(11)
    p_body.font.color.rgb = TEXT_LIGHT


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # =========================================================================
    # SLIDE 1: Capa (Title Slide)
    # =========================================================================
    s1 = create_blank_slide(prs)

    # Brilho decorativo
    glow = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(4.5), Inches(0.08))
    glow.fill.solid()
    glow.fill.fore_color.rgb = ACCENT_BLUE
    glow.line.fill.background()

    tb = s1.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p_tag = tf.paragraphs[0]
    p_tag.text = "APRESENTAÇÃO EXECUTIVA & ARQUITETURA DE PRODUTO"
    p_tag.font.name = FONT_HEADING
    p_tag.font.size = Pt(13)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT_BLUE
    p_tag.space_after = Pt(12)

    p_main = tf.add_paragraph()
    p_main.text = "MiniHub Car"
    p_main.font.name = FONT_HEADING
    p_main.font.size = Pt(48)
    p_main.font.bold = True
    p_main.font.color.rgb = TEXT_WHITE
    p_main.space_after = Pt(8)

    p_sub = tf.add_paragraph()
    p_sub.text = "A Plataforma Canônica de Gestão de Acervo, Marketplace e Comunidade para Colecionadores de Miniaturas"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = TEXT_LIGHT
    p_sub.space_after = Pt(24)

    p_desc = tf.add_paragraph()
    p_desc.text = "Foco em modelos premium escala 1:64 (Mini GT, Hot Wheels, Matchbox, Pop Race, Kaido House e Almost Real) • Catálogo de +16.000 miniaturas • Gestão Física de Estantes • Marketplace Homologado"
    p_desc.font.name = FONT_BODY
    p_desc.font.size = Pt(13)
    p_desc.font.color.rgb = TEXT_MUTED

    # Rodapé capa
    footer = s1.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.5))
    tf_f = footer.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "MiniHub Car Ecosystem • Versão 3.0 MVP & Roadmap Fase 2"
    p_f.font.name = FONT_BODY
    p_f.font.size = Pt(10)
    p_f.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 2: O Problema e a Solução
    # =========================================================================
    s2 = create_blank_slide(prs)
    add_header(s2, "Visão Estratégica", "O Desafio do Mercado e a Solução MiniHub Car",
               "Como transformamos a desorganização de colecionadores e lojistas em valor real.")

    add_card(
        s2, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.8),
        title="O Cenário Anterior (Dores)",
        text="• Dispersão Total: Coleções anotadas em cadernos, grupos de WhatsApp e planilhas confusas.\n\n"
             "• Desconhecimento do Patrimônio: Colecionadores não sabem quanto investiram e qual o valor real de mercado do seu acervo.\n\n"
             "• Localização Perdida: Dificuldade física de encontrar em qual caixa, prateleira ou gaveta o modelo está guardado.\n\n"
             "• Mercado Informal e Inseguro: Golpes em grupos fechados, falta de padronização nos nomes das miniaturas e fretes abusivos individuais.",
        accent_color=RGBColor(239, 68, 68),
        badge="Problema do Mercado"
    )

    add_card(
        s2, Inches(6.9), Inches(2.0), Inches(5.6), Inches(4.8),
        title="A Solução MiniHub Car",
        text="• Catálogo Oficial Canônico: +16.000 modelos com fotos em alta definição, especificações reais e montadoras originais.\n\n"
             "• Gestão Física Inteligente: Mapeamento em árvore (Cômodo > Móvel > Expositor > Nicho) com busca reversa imediata.\n\n"
             "• Marketplace com Vendedor Homologado: Vendedores verificados com estoque comercial segregado e ofertas padronizadas.\n\n"
             "• 'Minha Garagem' (Consolidação): O comprador acumula compras de semanas e despacha tudo em um frete único.",
        accent_color=ACCENT_GREEN,
        badge="Nossa Solução"
    )

    # =========================================================================
    # SLIDE 3: O que o Colecionador pode Esperar
    # =========================================================================
    s3 = create_blank_slide(prs)
    add_header(s3, "Experiência do Usuário", "O que o Colecionador pode Esperar da Aplicação?",
               "Recursos desenvolvidos sob medida para elevar a experiência do colecionismo.")

    cards_c = [
        ("Catalogação em 1 Clique",
         "Ao adicionar uma miniatura, o colecionador apenas pesquisa o modelo oficial e herda fotos HD, ano, série, casting e montadora real pré-preenchidos.",
         ACCENT_BLUE),
        ("Portfólio & Métricas de ROI",
         "Acompanhamento financeiro completo: valor total investido, estimativa de mercado atualizada e gráficos por marca e montadora automotiva.",
         ACCENT_GREEN),
        ("Wishlist Ativa com Match",
         "Lista de desejos categorizada por urgência. O sistema cruza sua wishlist com o Marketplace e avisa quando uma peça for anunciada.",
         ACCENT_ORANGE),
        ("Listas Temáticas & Compartilhamento",
         "Crie listas públicas ou privadas (ex: 'Top 10 Skylines', 'Projeto Le Mans') e compartilhe links dinâmicos diretamente com outros colecionadores.",
         ACCENT_PURPLE),
        ("Importação & Exportação Total",
         "Zero aprisionamento: exporte seu acervo completo para Excel/CSV ou importe planilhas antigas com relatório inteligente de inconsistências.",
         ACCENT_BLUE),
        ("PWA Mobile no Bolso",
         "App nativo instalável no Android e iPhone. Navegue offline ou em feiras e convenções para nunca mais comprar uma peça repetida.",
         ACCENT_GREEN),
    ]

    for i, (ctitle, cdesc, col) in enumerate(cards_c):
        row = i // 3
        col_idx = i % 3
        x = Inches(0.8 + col_idx * 4.0)
        y = Inches(2.0 + row * 2.5)
        add_card(s3, x, y, Inches(3.7), Inches(2.3), ctitle, cdesc, accent_color=col)

    # =========================================================================
    # SLIDE 4: O que o Vendedor Homologado pode Esperar
    # =========================================================================
    s4 = create_blank_slide(prs)
    add_header(s4, "Ecossistema Comercial", "O que o Vendedor Homologado pode Esperar?",
               "Um canal de vendas estruturado, profissional e focado exclusivamente em diecast de alto padrão.")

    cards_v = [
        ("Selo de Vendedor Homologado",
         "Processo de credenciamento rigoroso. O selo confere autoridade instantânea e atrai os colecionadores mais exigentes da plataforma.",
         ACCENT_BLUE, "Credibilidade"),
        ("Ofertas Conectadas ao Catálogo",
         "Chega de cadastrar anúncios do zero com nomes errados: vincule sua oferta ao modelo oficial e atinja 100% dos colecionadores que buscam a peça.",
         ACCENT_GREEN, "Precisão"),
        ("Gestão de Estoque Comercial",
         "Módulo dedicado (/seller) para precificação, quantidade, fotos reais da peça física e notas do blister, segregado do acervo pessoal.",
         ACCENT_PURPLE, "Controle"),
        ("Minha Garagem (Consolidação de Frete)",
         "O grande diferencial: o comprador adquire peças ao longo de 30 a 60 dias e solicita o envio em lote único. Aumenta o ticket médio do lojista.",
         ACCENT_ORANGE, "Logística & Economia"),
    ]

    for i, (ctitle, cdesc, col, badge) in enumerate(cards_v):
        row = i // 2
        col_idx = i % 2
        x = Inches(0.8 + col_idx * 6.0)
        y = Inches(2.0 + row * 2.5)
        add_card(s4, x, y, Inches(5.7), Inches(2.3), ctitle, cdesc, accent_color=col, badge=badge)

    # =========================================================================
    # SLIDE 5: Gestão de Localização Física
    # =========================================================================
    s5 = create_blank_slide(prs)
    add_header(s5, "Módulo Exclusivo", "Gestão Avançada de Localização Física do Acervo",
               "Organização física profissional para quem tem de dezenas a milhares de miniaturas.")

    add_card(
        s5, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.8),
        title="Hierarquia Física em Árvore",
        text="A plataforma replica perfeitamente o ambiente real do colecionador:\n\n"
             "🏛️ 1. Ambiente / Cômodo (ex: Escritório, Sala de Coleção)\n"
             "  └─ 🪟 2. Móvel / Expositor (ex: Estante de Vidro 01, Armário)\n"
             "      └─ 📐 3. Prateleira / Módulo (ex: Prateleira 2)\n"
             "          └─ 🏷️ 4. Nicho / Gaveta (ex: Vaga 14, Caixa Organizadora B)\n\n"
             "• O colecionador sabe exatamente onde cada modelo está, seja em exposição ou armazenado na caixa protetora.",
        accent_color=ACCENT_BLUE,
        badge="Estrutura Organizacional"
    )

    add_card(
        s5, Inches(6.9), Inches(2.0), Inches(5.6), Inches(4.8),
        title="Recursos Operacionais de Armazenamento",
        text="• Busca Reversa Imediata:\n"
             "  Digite 'Koenigsegg Gemera' e a tela indica imediatamente:\n"
             "  👉 Escritório > Estante Principal > Prateleira 2 > Nicho 14.\n\n"
             "• Gestão de Capacidade vs. Ocupação:\n"
             "  Visualize quantas vagas restam em cada expositor de acrílico (ex: 48 ocupadas de 60).\n\n"
             "• Estado de Conservação Física:\n"
             "  Classificação de cada unidade como Lacrada (Blister MINT), Solta (Loose), Com caixa original ou Customizada.",
        accent_color=ACCENT_ORANGE,
        badge="Controle Prático"
    )

    # =========================================================================
    # SLIDE 6: Métricas e Números do Catálogo Oficial
    # =========================================================================
    s6 = create_blank_slide(prs)
    add_header(s6, "Dados Reais", "Métricas do Catálogo Oficial MiniHub Car",
               "Uma das maiores bases de dados estruturadas de miniaturas automotivas da América Latina.")

    metrics = [
        ("Hot Wheels", "12.038+", "1968 a 2026 • Mainline, STH, TH, Premium e Séries Históricas", ACCENT_ORANGE),
        ("Matchbox", "1.322+", "2023 a 2026 atuais • Base histórica de +10.900 (1953 a 2027) • 1.248 fotos HD", ACCENT_BLUE),
        ("Mini GT", "1.940+", "MGT00001 a MGT01100+ • Kaido House, Mini GT Brasil e Especiais de Feiras", ACCENT_GREEN),
        ("Pop Race", "490+", "7 Coleções Completas (Regular, Enigma, Event, Dark Chrome, Blind Box)", ACCENT_PURPLE),
        ("Almost Real", "400+", "Miniaturas Ultra-Premium nas escalas 1:64, 1:43 e 1:18", ACCENT_BLUE),
        ("TOTAL", "+16.190", "Modelos cadastrados com fotos reais, montadoras e códigos oficiais", ACCENT_GREEN),
    ]

    for i, (mtitle, mval, msub, col) in enumerate(metrics):
        row = i // 3
        col_idx = i % 3
        x = Inches(0.8 + col_idx * 4.0)
        y = Inches(2.0 + row * 2.5)

        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.7), Inches(2.3))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.2)

        tb = s6.shapes.add_textbox(x + Inches(0.25), y + Inches(0.15), Inches(3.2), Inches(2.0))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = mtitle.upper()
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = col

        p_v = tf.add_paragraph()
        p_v.text = mval
        p_v.font.name = FONT_HEADING
        p_v.font.size = Pt(36)
        p_v.font.bold = True
        p_v.font.color.rgb = TEXT_WHITE
        p_v.space_after = Pt(4)

        p_s = tf.add_paragraph()
        p_s.text = msub
        p_s.font.name = FONT_BODY
        p_s.font.size = Pt(9.5)
        p_s.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 7: Todos os Módulos Desenvolvidos
    # =========================================================================
    s7 = create_blank_slide(prs)
    add_header(s7, "Visão Geral do Produto", "Panorama dos Módulos Desenvolvidos",
               "Arquitetura modular em microdomínios altamente integrados e tipados ponta a ponta.")

    modulos = [
        ("Catálogo Canônico", "Marca, Montadora real, Casting, Variação e identificadores oficiais (MAN, SKU, Toy#)."),
        ("Minha Coleção", "Gestão de itens físicos, fotos personalizadas, filtros por montadora e condição de conservação."),
        ("Localizações Físicas", "Hierarquia em árvore de cômodos, estantes e gavetas com controle de capacidade e busca reversa."),
        ("Wishlist & Listas", "Lista de desejos com níveis de prioridade e listas temáticas compartilháveis via link público."),
        ("Aquisições & Vendas", "Registro do histórico financeiro, cálculo de lucro apurado, ROI e evolução do patrimônio."),
        ("Mídia & Fotos HD", "Armazenamento e conversão otimizada de imagens JPEG com fundo branco e resolução máxima."),
        ("Marketplace & Ofertas", "Vitrine de produtos com vinculação estrita ao catálogo e ambiente para Vendedor Homologado."),
        ("Minha Garagem (Garage)", "Consolidação de frete inteligente: acumule compras de várias semanas e despache em lote único."),
        ("Governança & Curadoria", "Fluxo de sugestão de novas miniaturas por usuários e painel de aprovação/rejeição para Admins."),
        ("Sustentabilidade & Apoio", "Módulo de apoio comunitário via PIX com ciclo de 20 a 20 e transparência total de metas."),
    ]

    for i, (mtitle, mdesc) in enumerate(modulos):
        row = i // 5
        col_idx = i % 5
        x = Inches(0.8 + col_idx * 2.4)
        y = Inches(2.0 + row * 2.5)

        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.25), Inches(2.3))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tb = s7.shapes.add_textbox(x + Inches(0.12), y + Inches(0.12), Inches(2.0), Inches(2.0))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = mtitle
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = ACCENT_BLUE
        p_t.space_after = Pt(4)

        p_d = tf.add_paragraph()
        p_d.text = mdesc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 8: Roadmap Previsto para o Futuro
    # =========================================================================
    s8 = create_blank_slide(prs)
    add_header(s8, "Visão de Futuro", "Roadmap Estratégico (Fase 2 em Diante)",
               "Evolução contínua baseada nos 61 documentos de especificação técnica da Fase 2.")

    add_card(
        s8, Inches(0.8), Inches(2.0), Inches(3.7), Inches(4.8),
        title="Fase 1: Concluída",
        text="✅ Catálogo canônico unificado (+16k itens)\n\n"
             "✅ Coleção pessoal com fotos e estado\n\n"
             "✅ Gestão física de localizações e vagas\n\n"
             "✅ Wishlist inteligente e listas temáticas\n\n"
             "✅ Portfólio financeiro e histórico de compras\n\n"
             "✅ PWA Mobile instalável no celular\n\n"
             "✅ Deploy VPS seguro e isolado",
        accent_color=ACCENT_GREEN,
        badge="Fundação Sólida"
    )

    add_card(
        s8, Inches(4.8), Inches(2.0), Inches(3.7), Inches(4.8),
        title="Fase 2: Próximos Passos",
        text="🔄 Checkout Integrado com Custódia (Escrow):\n"
             "   Pagamento PIX/Cartão com liberação ao vendedor pós-recebimento.\n\n"
             "🔄 Cálculo de Frete Automático:\n"
             "   Integração direta com Correios / Melhor Envio para gerar etiquetas.\n\n"
             "🔄 Módulo Operacional de Rifas (Doc 10.20.9):\n"
             "   Gestão segura de cotas com sorteio pela Loteria Federal.\n\n"
             "🔄 Trade Hub:\n"
             "   Intermediação de trocas seguras entre colecionadores.",
        accent_color=ACCENT_ORANGE,
        badge="Transacional & Logística"
    )

    add_card(
        s8, Inches(8.8), Inches(2.0), Inches(3.7), Inches(4.8),
        title="Fase 3: Inovação & Escala",
        text="🚀 Scanner por Inteligência Artificial:\n"
             "   Aponte a câmera do celular para a miniatura solta e a IA identifica o modelo e variantes no catálogo.\n\n"
             "🚀 Índice de Preço Médio em Tempo Real:\n"
             "   Algoritmo que analisa vendas históricas e cotações globais (eBay / Brasil).\n\n"
             "🚀 Gamificação & Badges:\n"
             "   Pontuação de raridade do acervo e medalhas de reputação na comunidade.",
        accent_color=ACCENT_PURPLE,
        badge="Futuro & IA"
    )

    # =========================================================================
    # SLIDE 9: Arquitetura Tecnológica e Infraestrutura
    # =========================================================================
    s9 = create_blank_slide(prs)
    add_header(s9, "Engenharia de Software", "Arquitetura Tecnológica & Infraestrutura de Produção",
               "Tecnologias modernas, tipagem estrita e operação em contêineres sem conflito de portas.")

    add_card(
        s9, Inches(0.8), Inches(2.0), Inches(3.7), Inches(4.8),
        title="Backend & Dados",
        text="• Node.js + TypeScript (Strict Mode)\n\n"
             "• Fastify 5: Servidor HTTP de altíssima performance e baixa latência\n\n"
             "• PostgreSQL 16: Banco relacional robusto com integridade referencial estrita\n\n"
             "• Drizzle ORM: Mapeamento de dados totalmente tipado com migrations versionadas\n\n"
             "• Zod: Validação estrita de contratos de entrada e saída",
        accent_color=ACCENT_BLUE,
        badge="API RESTful"
    )

    add_card(
        s9, Inches(4.8), Inches(2.0), Inches(3.7), Inches(4.8),
        title="Frontend & Experiência",
        text="• Next.js 15 (App Router) + React 19\n\n"
             "• PWA (Progressive Web App): Manifesto e Service Worker para rodar como app nativo\n\n"
             "• Tailwind CSS: Design tokens automotivos escuros com alta legibilidade\n\n"
             "• TanStack Query: Cache e sincronização de dados remotos sem recarregamentos\n\n"
             "• Lucide Icons: Padronização visual moderna",
        accent_color=ACCENT_GREEN,
        badge="Interface Web & Mobile"
    )

    add_card(
        s9, Inches(8.8), Inches(2.0), Inches(3.7), Inches(4.8),
        title="Infraestrutura VPS",
        text="• Hostinger VPS (Ubuntu 24.04)\n\n"
             "• Coexistência Pacífica: Roda lado a lado com o AtivoArbóreo sem conflito:\n"
             "  - Postgres: porta 5434\n"
             "  - API: porta 3334\n"
             "  - Web: porta 3002\n\n"
             "• Proxy Reverso Nginx + SSL Certbot (HTTPS/2)\n\n"
             "• Volumes Docker montando o catálogo estático com cache de imagens de 30 dias",
        accent_color=ACCENT_ORANGE,
        badge="Deploy & Containers"
    )

    # =========================================================================
    # SLIDE 10: Conclusão & Encerramento
    # =========================================================================
    s10 = create_blank_slide(prs)

    add_card(
        s10, Inches(1.5), Inches(1.5), Inches(10.3), Inches(4.5),
        title="MiniHub Car: O Padrão Ouro para o Colecionismo de Miniaturas",
        text="Ao unir catálogo canônico em escala industrial, controle físico milimétrico de acervo, "
             "gestão financeira de valorização e um marketplace com consolidação de frete (Minha Garagem), "
             "o MiniHub Car se estabelece como a ferramenta definitiva para o colecionador apaixonado e o comerciante sério.\n\n"
             "✨ Catálogo pronto com +16.000 modelos e fotos em alta definição.\n"
             "✨ Pronta para publicação imediata na Hostinger VPS.\n"
             "✨ Base sólida para liderar o mercado de colecionismo automotivo no Brasil.",
        accent_color=ACCENT_BLUE,
        badge="Visão de Futuro"
    )

    footer10 = s10.shapes.add_textbox(Inches(1.5), Inches(6.3), Inches(10.3), Inches(0.6))
    tf_f10 = footer10.text_frame
    p_f10 = tf_f10.paragraphs[0]
    p_f10.text = "MiniHub Car • Obrigado! • Contato: admin@minihubcar.com.br • www.minihubcar.com.br"
    p_f10.alignment = PP_ALIGN.CENTER
    p_f10.font.name = FONT_HEADING
    p_f10.font.size = Pt(12)
    p_f10.font.color.rgb = ACCENT_BLUE

    # Salva apresentação
    output_path = Path("c:/Projetos/minihubcar/apresentacao_minihubcar.pptx")
    prs.save(str(output_path))
    print(f"✅ Apresentação salva com sucesso em: {output_path.resolve()}")


if __name__ == "__main__":
    build_presentation()
